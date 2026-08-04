"""Leest een teruggehaalde Excel-export en levert de verschillen op als een lijst
Wijziging (docs/DESIGN.md §4.4) -- de Drive-sheet is een bewerkingsoppervlak dat de
commissie rechtstreeks aanpast, dus terug-importeren moet een SAMENVOEGING zijn, geen
overschrijving.

Koppelvolgorde per rij, stop bij de eerste die lukt:
1. '_meta'-blad aanwezig en project_id klopt -> koppel op les_id.
2. Kolom 'Code' aanwezig -> koppel op id-prefix (alleen als uniek binnen het project).
3. Terugval op (dag, maand, begin_tijd) -- alleen als dat combinatie uniek is binnen het
   project. Er is BEWUST geen jaar in deze vergelijking: de Datum-kolom wordt als tekst
   geëxporteerd (bv. 'woensdag 22 apr', zie excel_export.py) en Excel/Sheets voegt daar bij
   het terugvertalen geen jaar aan toe -- exact dezelfde reden waarom forms_import.py's
   terugvalpad ook zonder jaar matcht.
Een rij die geen van drieën lukt wordt gerapporteerd, nooit geraden (Google Sheets kan
datums/tijden herformatteren, dus vertrouw niet blind op stap 3)."""
from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path

import openpyxl

from ..domain.formatting import parse_nl_maand
from ..domain.names import is_exact, stel_voor
from ..model.entities import Les
from ..model.project import Project
from .types import NaamProbleem, Wijziging


def lees_sheet(
    pad: str | Path, project: Project
) -> tuple[list[Wijziging], list[NaamProbleem], list[str]]:
    """Geeft (wijzigingen, naamproblemen, niet-koppelbare rijen) terug. De derde waarde is
    een uitbreiding op docs/DESIGN.md's signatuur (die liet in het midden waar
    niet-koppelbare rijen heen moeten) -- zie docs/BESLISSINGEN.md fase 7."""
    wb = openpyxl.load_workbook(str(pad), data_only=True)
    if "Planning" not in wb.sheetnames:
        raise ValueError("Geen blad 'Planning' gevonden -- is dit een geëxporteerd bestand?")
    ws = wb["Planning"]

    header = [_cel_tekst(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    kolom_index = {naam: i for i, naam in enumerate(header) if naam}

    tijd_col = kolom_index.get("Tijd")
    info_col = kolom_index.get("Info")
    if tijd_col is None or info_col is None:
        raise ValueError(
            "Kolommen 'Tijd' en/of 'Info' niet gevonden -- is dit een geëxporteerd "
            "planningsbestand?"
        )
    code_col = kolom_index.get("Code")
    datum_col = kolom_index.get("Datum")
    lesgever_cols = list(range(tijd_col + 1, info_col))

    meta_les_id_per_rij = _lees_meta(wb, project)

    lessen_by_id = {les.id: les for les in project.lessen}
    lessen_by_dag_maand_tijd: dict[tuple[int, int, time], list[Les]] = {}
    for les in project.lessen:
        sleutel = (les.datum.day, les.datum.month, les.begin_tijd)
        lessen_by_dag_maand_tijd.setdefault(sleutel, []).append(les)

    lesgever_naam_by_id = {lg.id: lg.naam for lg in project.lesgevers}
    naam_naar_lesgever = {lg.naam: lg for lg in project.lesgevers}

    wijzigingen: list[Wijziging] = []
    naamproblemen: list[NaamProbleem] = []
    niet_gekoppeld: list[str] = []

    for rijnummer, rij in enumerate(ws.iter_rows(min_row=2), start=2):
        waarden = [c.value for c in rij]
        if all(w in (None, "") for w in waarden):
            continue

        code_waarde = _cel_tekst(waarden[code_col]) if code_col is not None else ""
        les = None

        if meta_les_id_per_rij is not None:
            les_id = meta_les_id_per_rij.get(rijnummer)
            if les_id is not None:
                les = lessen_by_id.get(les_id)

        if les is None and code_waarde:
            kandidaten = [l for l in project.lessen if l.id.startswith(code_waarde)]
            if len(kandidaten) == 1:
                les = kandidaten[0]

        if les is None:
            dag_maand = _parse_dag_maand_cel(waarden[datum_col]) if datum_col is not None else None
            tijd_waarde = _parse_tijd_cel(waarden[tijd_col])
            if dag_maand is not None and tijd_waarde is not None:
                sleutel = (dag_maand[0], dag_maand[1], tijd_waarde)
                kandidaten = lessen_by_dag_maand_tijd.get(sleutel, [])
                if len(kandidaten) == 1:
                    les = kandidaten[0]

        if les is None:
            niet_gekoppeld.append(f"Rij {rijnummer} (code {code_waarde or '-'})")
            continue

        rij_wijzigingen, rij_naamproblemen = _vergelijk_rij(
            les, waarden, lesgever_cols, info_col, naam_naar_lesgever,
            lesgever_naam_by_id, project,
        )
        wijzigingen.extend(rij_wijzigingen)
        naamproblemen.extend(rij_naamproblemen)

    return wijzigingen, naamproblemen, niet_gekoppeld


def _vergelijk_rij(
    les: Les, waarden: list, lesgever_cols: list[int], info_col: int,
    naam_naar_lesgever: dict, lesgever_naam_by_id: dict, project: Project,
) -> tuple[list[Wijziging], list[NaamProbleem]]:
    wijzigingen: list[Wijziging] = []
    naamproblemen: list[NaamProbleem] = []

    ruwe_namen = [_cel_tekst(waarden[c]) for c in lesgever_cols if _cel_tekst(waarden[c])]
    is_geen_les = any(n.lower().startswith("geen les") for n in ruwe_namen)
    was_vervallen = les.status == "vervallen"

    if is_geen_les and not was_vervallen:
        reden = _parse_geen_les_reden(ruwe_namen)
        wijzigingen.append(
            Wijziging(
                les_id=les.id, soort="status", lesgever_id=None,
                oud="Gaat door",
                nieuw=f"Vervalt — {reden}" if reden else "Vervalt",
                zekerheid="exact",
            )
        )
    elif not is_geen_les and was_vervallen:
        oude_tekst = f"Vervalt — {les.vervallen_reden}" if les.vervallen_reden else "Vervalt"
        wijzigingen.append(
            Wijziging(
                les_id=les.id, soort="status", lesgever_id=None,
                oud=oude_tekst, nieuw="Gaat door", zekerheid="exact",
            )
        )

    if not is_geen_les:
        huidige_ids = {tw.lesgever_id for tw in les.toewijzingen}
        nieuwe_ids: set[str] = set()
        for naam in ruwe_namen:
            lesgever = naam_naar_lesgever.get(naam)
            if lesgever is None:
                voorstellen = stel_voor(naam, project.lesgevers)
                if voorstellen and is_exact(voorstellen[0][1]):
                    lesgever = voorstellen[0][0]
                else:
                    naamproblemen.append(
                        NaamProbleem(
                            ruwe_naam=naam,
                            voorstellen=[(lg.id, score) for lg, score in voorstellen[:5]],
                            les_id=les.id,
                        )
                    )
                    continue
            nieuwe_ids.add(lesgever.id)

        for lesgever_id in nieuwe_ids - huidige_ids:
            wijzigingen.append(
                Wijziging(
                    les_id=les.id, soort="toegevoegd", lesgever_id=lesgever_id,
                    oud="", nieuw=lesgever_naam_by_id.get(lesgever_id, "?"), zekerheid="exact",
                )
            )
        for lesgever_id in huidige_ids - nieuwe_ids:
            wijzigingen.append(
                Wijziging(
                    les_id=les.id, soort="verwijderd", lesgever_id=lesgever_id,
                    oud=lesgever_naam_by_id.get(lesgever_id, "?"), nieuw="", zekerheid="exact",
                )
            )

    sheet_info = _cel_tekst(waarden[info_col])
    huidige_info = les.titel or ""
    if sheet_info != huidige_info:
        wijzigingen.append(
            Wijziging(
                les_id=les.id, soort="titel", lesgever_id=None,
                oud=huidige_info, nieuw=sheet_info, zekerheid="exact",
            )
        )

    return wijzigingen, naamproblemen


def _cel_tekst(waarde) -> str:
    if waarde is None:
        return ""
    return str(waarde).strip()


def _parse_geen_les_reden(ruwe_namen: list[str]) -> str:
    for naam in ruwe_namen:
        if naam.lower().startswith("geen les"):
            rest = naam[len("geen les") :].strip()
            for scheiding in ("—", "-", ":"):
                if rest.startswith(scheiding):
                    return rest[len(scheiding) :].strip()
            return rest
    return ""


def _parse_dag_maand_cel(waarde) -> tuple[int, int] | None:
    """Geeft (dag, maand) terug zonder jaar -- zie de moduledocstring voor waarom."""
    if waarde is None:
        return None
    if isinstance(waarde, (datetime, date)):
        return (waarde.day, waarde.month)
    tekst = str(waarde).strip()
    delen = tekst.split()
    if len(delen) < 2:
        return None
    try:
        dag = int(delen[-2])
        maand = parse_nl_maand(delen[-1])
    except (ValueError, KeyError):
        return None
    return (dag, maand)


def _parse_tijd_cel(waarde) -> time | None:
    if waarde is None:
        return None
    if isinstance(waarde, datetime):
        return waarde.time()
    if isinstance(waarde, time):
        return waarde
    tekst = str(waarde).strip()
    eerste = tekst.split("-")[0].strip()
    try:
        uur, minuut = eerste.split(":")
        return time(int(uur), int(minuut))
    except ValueError:
        return None


def _lees_meta(wb, project: Project) -> dict[int, str] | None:
    if "_meta" not in wb.sheetnames:
        return None
    meta = wb["_meta"]
    project_id = None
    les_id_per_rij: dict[int, str] = {}
    modus = "kop"
    for row in meta.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        if row[0] == "project_id":
            project_id = row[1]
        elif row[0] == "les_id":
            modus = "tabel"
        elif modus == "tabel":
            les_id, rij_nummer = row[0], row[1]
            if les_id and rij_nummer:
                les_id_per_rij[int(rij_nummer)] = str(les_id)

    if project_id != project.id:
        return None
    return les_id_per_rij
