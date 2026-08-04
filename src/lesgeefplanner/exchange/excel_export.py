"""Exporteert de planning naar Excel. Zie docs/DESIGN.md §4.3.

De opmaak is bewust ONGEWIJZIGD ten opzichte van de oude src/export.py -- de commissie
leest dit blad in Google Drive, dus visuele verrassingen zijn ongewenst. Twee dingen zijn
toegevoegd t.o.v. de oude versie:

1. Een smalle 'Code'-kolom (eerste 8 tekens van het les-id) -- NIET verborgen, want een
   verborgen kolom overleeft de reis door Google Sheets niet betrouwbaar en mensen
   verwijderen wat ze niet begrijpen minder snel dan wat ze niet zien.
2. Een verborgen '_meta'-blad met project-id en een les_id-tabel, voor betrouwbare koppeling
   bij het terug-importeren (exchange/excel_import.py). Als iemand dat blad weggooit, valt
   import terug op de Code-kolom en daarna op (datum, begin_tijd).

`locale` wordt niet meer gebruikt -- Nederlandse dag-/maandnamen komen uit
domain/formatting.py."""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import xlsxwriter

from ..domain.formatting import MAANDEN_NL_KORT, format_dag, format_tijdvak
from ..model.entities import Les, les_seizoen_id
from ..model.project import Project

CODE_LENGTE = 8
COL_CODE = 0
COL_SEIZOEN = 1
COL_WEEK = 2
COL_DATUM = 3
COL_TIJD = 4
COL_LESGEVERS_START = 5
HEADER_ROWS = 1


def _lesgever_kolommen(project: Project) -> int:
    """Aantal lesgever-kolommen: minimaal het geconfigureerde maximum, maar nooit minder
    dan het grootste aantal toewijzingen dat een les daadwerkelijk heeft."""
    max_toegewezen = max((len(les.toewijzingen) for les in project.lessen), default=0)
    return max(project.solver_config.lesgever_maximum, max_toegewezen, 1)


def _col_letter(col: int) -> str:
    letters = ""
    col += 1
    while col > 0:
        col, rest = divmod(col - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def _week_nummer(d: date) -> int:
    return d.isocalendar()[1]


def _format_datum_lang_kort_maand(d: date) -> str:
    """bv. 'woensdag 03 sep' -- zelfde formaat als de oude export (%A %d %b onder nl_NL)."""
    return f"{format_dag(d, kort=False)} {d.day:02d} {MAANDEN_NL_KORT[d.month - 1]}"


def export_planning(project: Project, excel_pad: str | Path) -> None:
    workbook = xlsxwriter.Workbook(str(excel_pad))
    worksheet = workbook.add_worksheet("Planning")
    lesgever_kolommen = _lesgever_kolommen(project)
    col_info = COL_LESGEVERS_START + lesgever_kolommen

    seizoen_naam_by_id = {s.id: s.naam for s in project.seizoenen}
    lesgever_naam_by_id = {lg.id: lg.naam for lg in project.lesgevers}
    lessen = sorted(project.lessen, key=lambda l: (l.datum, l.begin_tijd))

    header_format = workbook.add_format({"align": "center", "bold": True, "top": 2})
    worksheet.write_string(0, COL_CODE, "Code", header_format)
    worksheet.write_string(0, COL_SEIZOEN, "Seizoen", header_format)
    worksheet.write_string(0, COL_WEEK, "Week", header_format)
    worksheet.write_string(0, COL_DATUM, "Datum", header_format)
    worksheet.write_string(0, COL_TIJD, "Tijd", header_format)
    lesgevers_range = (
        f"{_col_letter(COL_LESGEVERS_START)}1:"
        f"{_col_letter(COL_LESGEVERS_START + lesgever_kolommen - 1)}1"
    )
    worksheet.merge_range(lesgevers_range, "Lesgevers", header_format)
    worksheet.write_string(0, col_info, "Info", header_format)

    seizoen_waarden = [seizoen_naam_by_id.get(les_seizoen_id(les), "") for les in lessen]
    week_waarden = [_week_nummer(les.datum) for les in lessen]
    seizoen_ranges = _vind_bereiken(seizoen_waarden)
    week_ranges = _vind_bereiken(week_waarden)
    seizoen_start_rijen = {r[0] for r in seizoen_ranges}
    week_start_rijen = {r[0] for r in week_ranges}

    for row, les in enumerate(lessen):
        border_style = 0
        if row in week_start_rijen:
            border_style = 1
        if row in seizoen_start_rijen:
            border_style = 2
        _schrijf_les(
            worksheet, workbook, les, row + HEADER_ROWS, border_style,
            lesgever_kolommen, col_info, lesgever_naam_by_id,
        )

    for col in range(col_info + 1):
        row = len(lessen) + HEADER_ROWS
        worksheet.write_string(row, col, "", workbook.add_format({"top": 2}))

    _merge_kolom(
        worksheet, COL_SEIZOEN, seizoen_ranges, seizoen_waarden,
        workbook.add_format({"align": "center", "valign": "top", "bold": True, "top": 2}),
    )
    _merge_kolom(
        worksheet, COL_WEEK, week_ranges, week_waarden,
        workbook.add_format({"align": "center", "valign": "top", "bold": True, "top": 1}),
    )

    code_format = workbook.add_format({"font_size": 8, "font_color": "#999999"})
    worksheet.set_column(COL_CODE, COL_CODE, 3, code_format)
    worksheet.write_comment(
        0, COL_CODE, "Niet verwijderen -- nodig om wijzigingen terug in te lezen."
    )

    worksheet.autofit()
    workbook.close()

    _schrijf_meta_blad(str(excel_pad), project, lessen)


def _schrijf_les(
    worksheet, workbook, les: Les, row: int, border_style: int,
    lesgever_kolommen: int, col_info: int, lesgever_naam_by_id: dict[str, str],
) -> None:
    vervallen = les.status == "vervallen"
    fmt: dict = {"top": border_style}
    fmt["bg_color"] = "#DCE6F1" if _week_nummer(les.datum) % 2 == 0 else "#f0f0f0"
    if vervallen:
        fmt["font_color"] = "gray"
        fmt["italic"] = True

    format_ = workbook.add_format(fmt)
    l_border = workbook.add_format({"left": 1} | fmt)
    lr_border = workbook.add_format({"left": 1, "right": 1} | fmt)

    worksheet.write_string(row, COL_CODE, les.id[:CODE_LENGTE], format_)
    # COL_SEIZOEN wordt volledig door _merge_kolom() geschreven (elke rij zit in precies
    # één bereik), dus hier niets voorschrijven.
    worksheet.write_number(row, COL_WEEK, _week_nummer(les.datum), lr_border)
    worksheet.write_string(row, COL_DATUM, _format_datum_lang_kort_maand(les.datum), l_border)
    worksheet.write_string(row, COL_TIJD, format_tijdvak(les.begin_tijd, les.eind_tijd), format_)

    info_tekst = les.titel or ""
    worksheet.write_string(row, col_info, info_tekst, lr_border)

    if vervallen:
        links = _col_letter(COL_LESGEVERS_START)
        rechts = _col_letter(COL_LESGEVERS_START + lesgever_kolommen - 1)
        range_spec = f"{links}{row + 1}:{rechts}{row + 1}"
        tekst = f"Geen les — {les.vervallen_reden}" if les.vervallen_reden else "Geen les"
        merge_format = workbook.add_format({"align": "center", "left": 1} | fmt)
        worksheet.merge_range(range_spec, tekst, merge_format)
    else:
        namen = [lesgever_naam_by_id.get(tw.lesgever_id, "") for tw in les.toewijzingen]
        namen += [""] * (lesgever_kolommen - len(namen))
        for i, naam in enumerate(namen):
            col = COL_LESGEVERS_START + i
            worksheet.write(row, col, naam, l_border if i == 0 else format_)


def _vind_bereiken(waarden: list) -> list[tuple[int, int]]:
    grenzen = []
    vorige = object()
    for i, waarde in enumerate(waarden):
        if waarde != vorige:
            grenzen.append(i)
            vorige = waarde
    grenzen.append(len(waarden))
    return list(zip(grenzen[:-1], grenzen[1:]))


def _merge_kolom(worksheet, kolom: int, bereiken, waarden: list, format_) -> None:
    for top, bottom in bereiken:
        if bottom - top > 1:
            start_rij = top + HEADER_ROWS + 1
            eind_rij = bottom + HEADER_ROWS
            range_spec = f"{_col_letter(kolom)}{start_rij}:{_col_letter(kolom)}{eind_rij}"
            worksheet.merge_range(range_spec, waarden[top], format_)
        else:
            rij = top + HEADER_ROWS
            worksheet.write(rij, kolom, waarden[top], format_)


def _schrijf_meta_blad(excel_pad: str, project: Project, lessen: list[Les]) -> None:
    """Voegt een verborgen '_meta'-blad toe via openpyxl (xlsxwriter kan geen bestaand
    bestand aanvullen, dus dit is een tweede open/save-stap na het sluiten van xlsxwriter)."""
    import openpyxl

    wb = openpyxl.load_workbook(excel_pad)
    meta = wb.create_sheet("_meta")
    meta.append(["project_id", project.id])
    meta.append(["schema_version", project.schema_version])
    meta.append(["geexporteerd_op", datetime.now().isoformat()])
    meta.append([])
    meta.append(["les_id", "rij", "datum", "begin_tijd"])
    for rij, les in enumerate(lessen, start=HEADER_ROWS + 1):
        meta.append([les.id, rij, les.datum.isoformat(), les.begin_tijd.isoformat()])
    meta.sheet_state = "hidden"
    wb.save(excel_pad)
