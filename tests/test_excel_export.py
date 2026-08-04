from datetime import date, time
from pathlib import Path

import openpyxl

from lesgeefplanner.exchange.excel_export import export_planning
from lesgeefplanner.model import Lesgever, Project, Seizoen, Toewijzing
from lesgeefplanner.model.entities import Herkomst, Les


def _project() -> tuple[Project, Les, Les]:
    project = Project(naam="Test")
    seizoen = Seizoen(naam="Voorseizoen 1", begin=date(2026, 4, 19), eind=date(2026, 5, 10))
    project.seizoenen = [seizoen]
    lg = Lesgever(naam="Anne", ervaring_jaren=2)
    project.lesgevers = [lg]

    les1 = Les(
        datum=date(2026, 4, 22), begin_tijd=time(16, 0), eind_tijd=time(19, 0),
        seizoen_id=seizoen.id, herkomst=Herkomst(seizoen_id=seizoen.id, weekslot_index=0),
        toewijzingen=[Toewijzing(lesgever_id=lg.id, vast=True)],
    )
    les2 = Les(
        datum=date(2026, 4, 29), begin_tijd=time(16, 0), eind_tijd=time(19, 0),
        seizoen_id=seizoen.id, herkomst=Herkomst(seizoen_id=seizoen.id, weekslot_index=0),
        status="vervallen", vervallen_reden="Beka",
    )
    project.lessen = [les1, les2]
    return project, les1, les2


def test_export_schrijft_planning_en_meta_blad(tmp_path: Path):
    project, les1, les2 = _project()
    pad = tmp_path / "planning.xlsx"
    export_planning(project, pad)

    wb = openpyxl.load_workbook(pad)
    assert wb.sheetnames == ["Planning", "_meta"]
    assert wb["_meta"].sheet_state == "hidden"


def test_export_header_rij(tmp_path: Path):
    project, les1, les2 = _project()
    pad = tmp_path / "planning.xlsx"
    export_planning(project, pad)
    ws = openpyxl.load_workbook(pad)["Planning"]
    header = [c.value for c in ws[1]]
    assert header[0] == "Code"
    assert header[1] == "Seizoen"
    assert "Lesgevers" in header
    assert header[-1] == "Info"


def test_export_code_kolom_matcht_les_id_prefix(tmp_path: Path):
    project, les1, les2 = _project()
    pad = tmp_path / "planning.xlsx"
    export_planning(project, pad)
    ws = openpyxl.load_workbook(pad)["Planning"]
    codes = [ws.cell(row=r, column=1).value for r in (2, 3)]
    assert codes == [les1.id[:8], les2.id[:8]]


def test_export_meta_les_id_tabel(tmp_path: Path):
    project, les1, les2 = _project()
    pad = tmp_path / "planning.xlsx"
    export_planning(project, pad)
    meta = openpyxl.load_workbook(pad)["_meta"]

    gevonden = {}
    lees_tabel = False
    for row in meta.iter_rows(values_only=True):
        if row[0] == "les_id":
            lees_tabel = True
            continue
        if lees_tabel and row[0]:
            gevonden[row[0]] = row[1]

    assert gevonden[les1.id] == 2
    assert gevonden[les2.id] == 3


def test_export_vervallen_les_toont_reden(tmp_path: Path):
    project, les1, les2 = _project()
    pad = tmp_path / "planning.xlsx"
    export_planning(project, pad)
    ws = openpyxl.load_workbook(pad)["Planning"]
    cel = ws.cell(row=3, column=6).value
    assert cel is not None
    assert "Beka" in cel
    assert "—" in cel


def test_export_lesgever_naam_zichtbaar(tmp_path: Path):
    project, les1, les2 = _project()
    pad = tmp_path / "planning.xlsx"
    export_planning(project, pad)
    ws = openpyxl.load_workbook(pad)["Planning"]
    rij_waarden = [c.value for c in ws[2]]
    assert "Anne" in rij_waarden
